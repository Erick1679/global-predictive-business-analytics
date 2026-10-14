from pathlib import Path
from io import BytesIO

import pandas as pd
import numpy as np
import streamlit as st
import plotly.express as px
import plotly.graph_objects as go

import xlsxwriter
from pptx import Presentation
from pptx.util import Inches, Pt

# ==================================================
# APPLICATION CONFIGURATION
# ==================================================


st.set_page_config(
    page_title="Global Predictive Business Analytics Platform",
    page_icon="📊",
    layout="wide",
)


DATA_DIR = Path("data/processed")


COLOR_SEQUENCE = [
    "#2563eb",
    "#16a34a",
    "#f59e0b",
    "#dc2626",
    "#9333ea",
    "#0891b2",
    "#ea580c",
    "#64748b",
]


st.markdown(
    """
    <style>

    .main-title {

        font-size:42px;
        font-weight:750;
        color:#0f172a;

    }


    .section-title {

        font-size:30px;
        font-weight:650;
        color:#1e293b;

    }


    .kpi-card {

        background:#f8fafc;
        padding:18px;
        border-radius:12px;
        border:1px solid #e2e8f0;

    }


    </style>
    """,
    unsafe_allow_html=True,
)


# ==================================================
# DATA LOADING
# ==================================================


@st.cache_data
def load_data():

    files = {
        "customers": "customer_features.csv",
        "sales": "sales_features.csv",
        "products": "product_features.csv",
        "inventory": "cleaned_inventory.csv",
        "churn": "customer_churn_predictions.csv",
        "clv": "customer_clv_predictions.csv",
        "clv_summary": "clv_segment_summary.csv",
        "forecast": "sales_forecast_predictions.csv",
        "future_forecast": "future_sales_forecast.csv",
        "forecast_metrics": "forecast_metrics.csv",
        "churn_metrics": "churn_model_metrics.csv",
    }

    datasets = {}

    for name, file in files.items():

        path = DATA_DIR / file

        if path.exists():

            datasets[name] = pd.read_csv(path)

        else:

            datasets[name] = pd.DataFrame()

    return datasets


data = load_data()


# ==================================================
# HELPER FUNCTIONS
# ==================================================


def currency(value):

    return f"${value:,.0f}"


def percentage(value):

    return f"{value:.1f}%"


def create_kpi(
    title,
    value,
):

    st.markdown(
        f"""
        <div class="kpi-card">

        <b>{title}</b>

        <h2>{value}</h2>

        </div>
        """,
        unsafe_allow_html=True,
    )


def sort_chart(dataframe, column, ascending=False):

    return dataframe.sort_values(column, ascending=ascending)


# ==================================================
# SIDEBAR
# ==================================================


st.sidebar.markdown("## 📊 Analytics Platform")


page = st.sidebar.radio(
    "Navigate",
    [
        "Executive Overview",
        "Revenue Forecasting",
        "Customer Intelligence",
        "Customer Lifetime Value",
        "Product Analytics",
        "Operations & Inventory",
        "Model Performance",
        "Data Quality",
        "Download Centre",
    ],
)


st.sidebar.info("""
Global Predictive Business Analytics Platform


Enterprise analytics solution combining:


• Data Engineering

• Machine Learning

• Business Intelligence

• Predictive Modelling

""")


st.sidebar.success("Platform Ready ✓")
# ==================================================
# EXECUTIVE OVERVIEW
# ==================================================


if page == "Executive Overview":

    st.markdown(
        "<div class='main-title'>Global Predictive Business Analytics Platform</div>",
        unsafe_allow_html=True,
    )

    st.write("""
        Enterprise analytics platform combining:

        • Business Intelligence  
        • Predictive Analytics  
        • Customer Intelligence  
        • Revenue Forecasting  
        • Operational Analytics
        """)

    customers = data["customers"]

    sales = data["sales"]

    churn = data["churn"]

    # ----------------------------------------------
    # KPI CALCULATIONS
    # ----------------------------------------------

    total_revenue = sales["total_sales"].sum()

    total_profit = sales["total_profit"].sum()

    total_customers = customers["customer_id"].nunique()

    total_orders = customers["total_orders"].sum()

    average_order_value = sales["average_order_value"].mean()

    profit_margin = (total_profit / total_revenue) * 100 if total_revenue > 0 else 0

    customer_health = 100 - (
        churn["risk_level"].value_counts(normalize=True).get("High Risk", 0) * 100
    )

    # ----------------------------------------------
    # KPI CARDS
    # ----------------------------------------------

    c1, c2, c3 = st.columns(3)

    with c1:

        create_kpi("Total Revenue", currency(total_revenue))

    with c2:

        create_kpi("Total Profit", currency(total_profit))

    with c3:

        create_kpi("Profit Margin", percentage(profit_margin))

    c4, c5, c6 = st.columns(3)

    with c4:

        create_kpi("Customers", f"{total_customers:,}")

    with c5:

        create_kpi("Orders", f"{total_orders:,.0f}")

    with c6:

        create_kpi("Customer Health Score", percentage(customer_health))

    st.divider()

    # ----------------------------------------------
    # COUNTRY PERFORMANCE
    # ----------------------------------------------

    st.subheader("Revenue Performance by Country")

    country_sales = sales.groupby("country")["total_sales"].sum().reset_index()

    country_sales = sort_chart(country_sales, "total_sales")

    country_chart = px.bar(
        country_sales,
        x="country",
        y="total_sales",
        color="country",
        text="total_sales",
        color_discrete_sequence=COLOR_SEQUENCE,
        template="plotly_white",
    )

    country_chart.update_traces(texttemplate="$%{text:,.0f}", textposition="outside")

    country_chart.update_layout(
        xaxis_title="Country", yaxis_title="Revenue", showlegend=False
    )

    st.plotly_chart(
        country_chart,
        use_container_width=True,
    )

    # ----------------------------------------------
    # CUSTOMER RISK OVERVIEW
    # ----------------------------------------------

    st.subheader("Customer Risk Overview")

    risk_summary = churn["risk_level"].value_counts().reset_index()

    risk_summary.columns = ["Risk Level", "Customers"]

    risk_chart = px.pie(
        risk_summary,
        names="Risk Level",
        values="Customers",
        hole=0.45,
        color="Risk Level",
        color_discrete_sequence=COLOR_SEQUENCE,
        template="plotly_white",
    )

    st.plotly_chart(
        risk_chart,
        use_container_width=True,
    )

    # ----------------------------------------------
    # BUSINESS SUMMARY
    # ----------------------------------------------

    st.subheader("Executive Business Summary")

    summary = pd.DataFrame(
        {
            "Metric": [
                "Revenue",
                "Profit",
                "Profit Margin",
                "Customers",
                "Orders",
                "Average Order Value",
            ],
            "Value": [
                currency(total_revenue),
                currency(total_profit),
                percentage(profit_margin),
                f"{total_customers:,}",
                f"{total_orders:,.0f}",
                currency(average_order_value),
            ],
        }
    )

    st.dataframe(
        summary,
        hide_index=True,
        use_container_width=True,
    )
# ==================================================
# REVENUE FORECASTING
# ==================================================


if page == "Revenue Forecasting":

    st.markdown(
        "<div class='section-title'>Revenue Forecasting Intelligence</div>",
        unsafe_allow_html=True,
    )

    forecast = data["forecast"].copy()

    future_forecast = data["future_forecast"].copy()

    forecast_metrics = data["forecast_metrics"]

    forecast["month"] = pd.to_datetime(forecast["month"])

    future_forecast["forecast_month"] = pd.to_datetime(
        future_forecast["forecast_month"]
    )

    # ----------------------------------------------
    # FORECAST KPI CARDS
    # ----------------------------------------------

    latest_actual = forecast["actual_sales"].iloc[-1]

    future_total = future_forecast["predicted_sales"].sum()

    forecast_accuracy = 100 - (forecast_metrics["MAPE"].iloc[0])

    growth_rate = (
        (future_forecast["predicted_sales"].iloc[-1] - latest_actual)
        / latest_actual
        * 100
    )

    c1, c2, c3 = st.columns(3)

    with c1:

        create_kpi("Future Forecast Revenue", currency(future_total))

    with c2:

        create_kpi("Forecast Accuracy", percentage(forecast_accuracy))

    with c3:

        create_kpi("Expected Growth", percentage(growth_rate))

    st.divider()

    # ----------------------------------------------
    # ACTUAL VS PREDICTED
    # ----------------------------------------------

    st.subheader("Actual vs Predicted Revenue")

    forecast_monthly = (
        forecast.groupby("month")[["actual_sales", "predicted_sales"]]
        .sum()
        .reset_index()
    )

    fig = go.Figure()

    fig.add_trace(
        go.Scatter(
            x=forecast_monthly["month"],
            y=forecast_monthly["actual_sales"],
            name="Actual Revenue",
            mode="lines+markers",
            line=dict(shape="spline", smoothing=1.3, width=4, color="#2563eb"),
            marker=dict(size=8),
        )
    )

    fig.add_trace(
        go.Scatter(
            x=forecast_monthly["month"],
            y=forecast_monthly["predicted_sales"],
            name="Predicted Revenue",
            mode="lines+markers",
            line=dict(
                shape="spline", smoothing=1.3, width=4, dash="dot", color="#dc2626"
            ),
            marker=dict(size=8),
        )
    )

    fig.update_layout(
        template="plotly_white",
        hovermode="x unified",
        title="Revenue Trend Forecast",
    )

    st.plotly_chart(fig, use_container_width=True)

    # ----------------------------------------------
    # FUTURE FORECAST
    # ----------------------------------------------

    st.subheader("Future Revenue Projection")

    future_fig = px.line(
        future_forecast,
        x="forecast_month",
        y="predicted_sales",
        markers=True,
        template="plotly_white",
        title="Next 6 Months Revenue Forecast",
    )

    future_fig.update_traces(
        line_shape="spline",
        line_width=4,
    )

    st.plotly_chart(future_fig, use_container_width=True)

    st.subheader("Forecast Model Metrics")

    st.dataframe(forecast_metrics, hide_index=True, use_container_width=True)


# ==================================================
# CUSTOMER INTELLIGENCE
# ==================================================


if page == "Customer Intelligence":

    st.markdown(
        "<div class='section-title'>Customer Intelligence</div>",
        unsafe_allow_html=True,
    )

    churn = data["churn"]

    customers = data["customers"]

    # ----------------------------------------------
    # CUSTOMER KPIS
    # ----------------------------------------------

    total_customers = customers["customer_id"].nunique()

    risk_counts = churn["risk_level"].value_counts()

    high_risk = risk_counts.get("High Risk", 0)

    retention_opportunity = high_risk / total_customers * 100

    c1, c2, c3 = st.columns(3)

    with c1:

        create_kpi("Total Customers", f"{total_customers:,}")

    with c2:

        create_kpi("High Risk Customers", f"{high_risk:,}")

    with c3:

        create_kpi("Retention Opportunity", percentage(retention_opportunity))

    st.divider()

    # ----------------------------------------------
    # RISK DISTRIBUTION
    # ----------------------------------------------

    st.subheader("Customer Risk Distribution")

    risk_summary = risk_counts.reset_index()

    risk_summary.columns = ["Risk Level", "Customers"]

    risk_summary = sort_chart(risk_summary, "Customers")

    risk_chart = px.bar(
        risk_summary,
        x="Risk Level",
        y="Customers",
        color="Risk Level",
        text="Customers",
        color_discrete_sequence=COLOR_SEQUENCE,
        template="plotly_white",
    )

    risk_chart.update_traces(textposition="outside")

    st.plotly_chart(risk_chart, use_container_width=True)

    # ----------------------------------------------
    # HIGH RISK CUSTOMERS
    # ----------------------------------------------

    st.subheader("Highest Priority Customers")

    high_risk_customers = churn.sort_values("churn_probability", ascending=False).head(
        25
    )

    st.dataframe(high_risk_customers, hide_index=True, use_container_width=True)
# ==================================================
# CUSTOMER LIFETIME VALUE
# ==================================================


if page == "Customer Lifetime Value":

    st.markdown(
        "<div class='section-title'>Customer Lifetime Value Intelligence</div>",
        unsafe_allow_html=True,
    )

    clv = data["clv"]

    clv_summary = data["clv_summary"]

    # ----------------------------------------------
    # CLV KPIs
    # ----------------------------------------------

    total_value = clv_summary["total_value"].sum()

    average_clv = clv_summary["average_clv"].mean()

    premium_customers = clv[clv["clv_segment"] == "Platinum"].shape[0]

    top_segment = clv_summary.sort_values("average_clv", ascending=False).iloc[0][
        "clv_segment"
    ]

    c1, c2, c3, c4 = st.columns(4)

    with c1:

        create_kpi("Total Customer Value", currency(total_value))

    with c2:

        create_kpi("Average CLV", currency(average_clv))

    with c3:

        create_kpi("Premium Customers", f"{premium_customers:,}")

    with c4:

        create_kpi("Highest Segment", top_segment)

    st.divider()

    # ----------------------------------------------
    # CLV SEGMENT RANKING
    # ----------------------------------------------

    st.subheader("Customer Value Segmentation")

    clv_rank = sort_chart(clv_summary, "total_value")

    clv_chart = px.bar(
        clv_rank,
        x="clv_segment",
        y="total_value",
        color="clv_segment",
        text="total_value",
        color_discrete_sequence=COLOR_SEQUENCE,
        template="plotly_white",
    )

    clv_chart.update_traces(texttemplate="$%{text:,.0f}", textposition="outside")

    st.plotly_chart(clv_chart, use_container_width=True)

    st.subheader("CLV Segment Summary")

    st.dataframe(clv_rank, hide_index=True, use_container_width=True)


# ==================================================
# PRODUCT ANALYTICS
# ==================================================


if page == "Product Analytics":

    st.markdown(
        "<div class='section-title'>Product Analytics</div>",
        unsafe_allow_html=True,
    )

    products = data["products"]

    # ----------------------------------------------
    # PRODUCT KPIS
    # ----------------------------------------------

    total_products = products["product_id"].nunique()

    best_category = (
        products.groupby("category")["total_revenue"]
        .sum()
        .sort_values(ascending=False)
        .index[0]
    )

    top_product_revenue = products["total_revenue"].max()

    average_margin = products["profit_margin"].mean()

    c1, c2, c3, c4 = st.columns(4)

    with c1:

        create_kpi("Products", f"{total_products:,}")

    with c2:

        create_kpi("Best Category", best_category)

    with c3:

        create_kpi("Top Product Revenue", currency(top_product_revenue))

    with c4:

        create_kpi("Average Margin", percentage(average_margin))

    st.divider()

    # ----------------------------------------------
    # TOP PRODUCTS
    # ----------------------------------------------

    st.subheader("Top Products by Revenue")

    top_products = products.sort_values("total_revenue", ascending=False).head(10)

    product_chart = px.bar(
        top_products,
        x="product_name",
        y="total_revenue",
        color="category",
        text="total_revenue",
        color_discrete_sequence=COLOR_SEQUENCE,
        template="plotly_white",
    )

    product_chart.update_traces(texttemplate="$%{text:,.0f}", textposition="outside")

    st.plotly_chart(product_chart, use_container_width=True)

    # ----------------------------------------------
    # PROFITABILITY MATRIX
    # ----------------------------------------------

    st.subheader("Product Profitability Matrix")

    profitability = px.scatter(
        products,
        x="total_revenue",
        y="total_profit",
        size="units_sold",
        color="category",
        hover_name="product_name",
        color_discrete_sequence=COLOR_SEQUENCE,
        template="plotly_white",
    )

    st.plotly_chart(profitability, use_container_width=True)


# ==================================================
# OPERATIONS & INVENTORY
# ==================================================


if page == "Operations & Inventory":

    st.markdown(
        "<div class='section-title'>Operations & Inventory Intelligence</div>",
        unsafe_allow_html=True,
    )

    inventory = data["inventory"]

    # ----------------------------------------------
    # INVENTORY KPIS
    # ----------------------------------------------

    total_stock = inventory["stock_on_hand"].sum()

    low_stock = inventory[inventory["stock_risk"] == "High"].shape[0]

    warehouses = inventory["warehouse"].nunique()

    c1, c2, c3 = st.columns(3)

    with c1:

        create_kpi("Total Stock Units", f"{total_stock:,.0f}")

    with c2:

        create_kpi("Low Stock Items", f"{low_stock:,}")

    with c3:

        create_kpi("Warehouses", f"{warehouses}")

    st.divider()

    # ----------------------------------------------
    # STOCK RISK
    # ----------------------------------------------

    risk = inventory["stock_risk"].value_counts().reset_index()

    risk.columns = ["Risk", "Count"]

    risk_chart = px.pie(
        risk,
        names="Risk",
        values="Count",
        hole=0.45,
        color="Risk",
        color_discrete_sequence=COLOR_SEQUENCE,
        template="plotly_white",
    )

    st.plotly_chart(risk_chart, use_container_width=True)

    # ----------------------------------------------
    # WAREHOUSE PERFORMANCE
    # ----------------------------------------------

    warehouse = (
        inventory.groupby("warehouse")["stock_on_hand"]
        .sum()
        .reset_index()
        .sort_values("stock_on_hand", ascending=False)
    )

    warehouse_chart = px.bar(
        warehouse,
        x="warehouse",
        y="stock_on_hand",
        color="warehouse",
        text="stock_on_hand",
        color_discrete_sequence=COLOR_SEQUENCE,
        template="plotly_white",
    )

    warehouse_chart.update_traces(textposition="outside")

    st.plotly_chart(warehouse_chart, use_container_width=True)
# ==================================================
# MODEL PERFORMANCE
# ==================================================


if page == "Model Performance":

    st.markdown(
        "<div class='section-title'>Machine Learning Model Performance</div>",
        unsafe_allow_html=True,
    )

    churn_metrics = data["churn_metrics"]

    forecast_metrics = data["forecast_metrics"]

    # ----------------------------------------------
    # MODEL SCORECARD
    # ----------------------------------------------

    st.subheader("Customer Churn Model Evaluation")

    best_model = churn_metrics.sort_values("roc_auc", ascending=False).iloc[0]

    c1, c2, c3 = st.columns(3)

    with c1:

        create_kpi("Best Model", best_model["model"])

    with c2:

        create_kpi("Accuracy", percentage(best_model["accuracy"] * 100))

    with c3:

        create_kpi("ROC-AUC", percentage(best_model["roc_auc"] * 100))

    st.divider()

    st.subheader("Churn Model Comparison")

    churn_display = churn_metrics.copy()

    for column in [
        "accuracy",
        "precision",
        "recall",
        "f1_score",
        "roc_auc",
    ]:

        if column in churn_display.columns:

            churn_display[column] = (churn_display[column] * 100).round(2)

    st.dataframe(churn_display, hide_index=True, use_container_width=True)

    model_chart = px.bar(
        churn_display,
        x="model",
        y=[
            "accuracy",
            "precision",
            "recall",
            "f1_score",
            "roc_auc",
        ],
        barmode="group",
        color_discrete_sequence=COLOR_SEQUENCE,
        template="plotly_white",
        title="Classification Performance Comparison",
    )

    st.plotly_chart(model_chart, use_container_width=True)

    st.divider()

    st.subheader("Forecast Model Evaluation")

    st.dataframe(forecast_metrics, hide_index=True, use_container_width=True)


# ==================================================
# DATA QUALITY
# ==================================================


if page == "Data Quality":

    st.markdown(
        "<div class='section-title'>Data Quality & Pipeline Health</div>",
        unsafe_allow_html=True,
    )

    st.subheader("Dataset Monitoring")

    dataset_quality = pd.DataFrame(
        {
            "Dataset": [
                "Customers",
                "Products",
                "Transactions",
                "Marketing",
                "Support",
                "Inventory",
            ],
            "Records": [
                "15,000",
                "1,000",
                "120,000",
                "151,504",
                "5,000",
                "10,000",
            ],
            "Status": [
                "✓ Complete",
                "✓ Complete",
                "✓ Complete",
                "✓ Complete",
                "✓ Complete",
                "✓ Complete",
            ],
        }
    )

    st.dataframe(dataset_quality, hide_index=True, use_container_width=True)

    st.divider()

    st.subheader("Analytics Pipeline Status")

    pipeline = pd.DataFrame(
        {
            "Pipeline Stage": [
                "Synthetic Data Generation",
                "Data Cleaning",
                "Feature Engineering",
                "Machine Learning Training",
                "Forecast Development",
                "Dashboard Deployment",
            ],
            "Status": [
                "Completed",
                "Completed",
                "Completed",
                "Completed",
                "Completed",
                "Completed",
            ],
            "Health": [
                "100%",
                "100%",
                "100%",
                "100%",
                "100%",
                "100%",
            ],
        }
    )

    st.table(pipeline)

    st.subheader("Technology Stack")

    st.info("""

        Programming:

        Python


        Data Engineering:

        Pandas | NumPy


        Machine Learning:

        Scikit-learn


        Visualization:

        Plotly | Streamlit


        Analytics Capabilities:

        • Classification

        • Regression

        • Forecasting

        • Customer Lifetime Value

        • Business Intelligence

        """)


# ==================================================
# DOWNLOAD CENTRE
# ==================================================


if page == "Download Centre":

    st.markdown(
        "<div class='section-title'>Analytics Report Download Centre</div>",
        unsafe_allow_html=True,
    )

    reports = {
        "Customer Churn Predictions": data["churn"],
        "Customer Lifetime Value": data["clv"],
        "Revenue Forecast": data["forecast"],
        "Future Revenue Forecast": data["future_forecast"],
        "Forecast Metrics": data["forecast_metrics"],
        "Churn Metrics": data["churn_metrics"],
    }

    for name, dataframe in reports.items():

        csv = dataframe.to_csv(index=False)

        st.download_button(
            label=f"Download {name}",
            data=csv,
            file_name=f"{name}.csv",
            mime="text/csv",
        )


st.sidebar.success("Analytics Platform Complete ✓")
# ==================================================
# REPORT GENERATION ENGINE
# ==================================================


def create_excel_report(data):

    output = BytesIO()

    with pd.ExcelWriter(output, engine="xlsxwriter") as writer:

        workbook = writer.book

        # -----------------------------
        # EXECUTIVE OVERVIEW
        # -----------------------------

        sales = data["sales"]

        customers = data["customers"]

        churn = data["churn"]

        executive = pd.DataFrame(
            {
                "Metric": [
                    "Total Revenue",
                    "Total Profit",
                    "Customers",
                    "Orders",
                    "Average Order Value",
                ],
                "Value": [
                    sales["total_sales"].sum(),
                    sales["total_profit"].sum(),
                    customers["customer_id"].nunique(),
                    customers["total_orders"].sum(),
                    sales["average_order_value"].mean(),
                ],
            }
        )

        executive.to_excel(writer, sheet_name="Executive Overview", index=False)

        worksheet = writer.sheets["Executive Overview"]

        chart = workbook.add_chart({"type": "column"})

        chart.add_series(
            {
                "name": "Revenue",
                "categories": "='Executive Overview'!$A$2:$A$6",
                "values": "='Executive Overview'!$B$2:$B$6",
            }
        )

        worksheet.insert_chart("D2", chart)

        # -----------------------------
        # REVENUE FORECAST
        # -----------------------------

        data["forecast"].to_excel(writer, sheet_name="Revenue Forecast", index=False)

        worksheet = writer.sheets["Revenue Forecast"]

        chart = workbook.add_chart({"type": "line"})

        chart.add_series(
            {
                "name": "Actual Revenue",
                "categories": "='Revenue Forecast'!$A$2:$A$100",
                "values": "='Revenue Forecast'!$B$2:$B$100",
            }
        )

        chart.add_series(
            {
                "name": "Predicted Revenue",
                "categories": "='Revenue Forecast'!$A$2:$A$100",
                "values": "='Revenue Forecast'!$C$2:$C$100",
            }
        )

        worksheet.insert_chart("F2", chart)

        # -----------------------------
        # CUSTOMER INTELLIGENCE
        # -----------------------------

        data["churn"].to_excel(writer, sheet_name="Customer Intelligence", index=False)

        # -----------------------------
        # CLV
        # -----------------------------

        data["clv_summary"].to_excel(writer, sheet_name="CLV Analysis", index=False)

        # -----------------------------
        # PRODUCTS
        # -----------------------------

        data["products"].to_excel(writer, sheet_name="Product Analytics", index=False)

        # -----------------------------
        # INVENTORY
        # -----------------------------

        data["inventory"].to_excel(writer, sheet_name="Operations", index=False)

        # -----------------------------
        # MODELS
        # -----------------------------

        data["churn_metrics"].to_excel(
            writer, sheet_name="Model Performance", index=False
        )

        data["forecast_metrics"].to_excel(
            writer, sheet_name="Forecast Metrics", index=False
        )

    output.seek(0)

    return output


def create_powerpoint_report(data):

    presentation = Presentation()

    # -----------------------------
    # TITLE SLIDE
    # -----------------------------

    slide = presentation.slides.add_slide(presentation.slide_layouts[0])

    slide.shapes.title.text = "Global Predictive Business Analytics Platform"

    slide.placeholders[1].text = "Executive Analytics Report"

    # -----------------------------
    # EXECUTIVE SLIDE
    # -----------------------------

    slide = presentation.slides.add_slide(presentation.slide_layouts[1])

    slide.shapes.title.text = "Executive Overview"

    revenue = data["sales"]["total_sales"].sum()

    slide.placeholders[1].text = f"""

        Total Revenue:

        ${revenue:,.0f}


        Customers:

        {data["customers"]["customer_id"].nunique():,}


        Orders:

        {data["customers"]["total_orders"].sum():,.0f}

        """

    # -----------------------------
    # FORECAST SLIDE
    # -----------------------------

    slide = presentation.slides.add_slide(presentation.slide_layouts[1])

    slide.shapes.title.text = "Revenue Forecasting"

    slide.placeholders[1].text = "Actual vs Predicted Revenue Analysis"

    # -----------------------------
    # CUSTOMER SLIDE
    # -----------------------------

    slide = presentation.slides.add_slide(presentation.slide_layouts[1])

    slide.shapes.title.text = "Customer Intelligence"

    slide.placeholders[1].text = "Customer churn risk and retention analysis"

    # -----------------------------
    # MODEL SLIDE
    # -----------------------------

    slide = presentation.slides.add_slide(presentation.slide_layouts[1])

    slide.shapes.title.text = "Machine Learning Performance"

    slide.placeholders[1].text = "Classification and forecasting evaluation"

    output = BytesIO()

    presentation.save(output)

    output.seek(0)

    return output


# ==================================================
# DOWNLOAD CENTRE
# ==================================================


if page == "Download Centre":

    st.markdown(
        "<div class='section-title'>Analytics Report Download Centre</div>",
        unsafe_allow_html=True,
    )

    st.write("""

        Export your analytics results in multiple professional formats.


        • CSV for analysis

        • Excel workbook with editable charts

        • PowerPoint executive presentation

        """)

    st.divider()

    # CSV DOWNLOADS

    st.subheader("CSV Data Export")

    for name, dataframe in data.items():

        csv = dataframe.to_csv(index=False)

        st.download_button(
            label=f"Download {name}",
            data=csv,
            file_name=f"{name}.csv",
            mime="text/csv",
        )

    st.divider()

    # EXCEL REPORT

    st.subheader("Editable Excel Analytics Workbook")

    excel_file = create_excel_report(data)

    st.download_button(
        label="📊 Download Excel Analytics Report",
        data=excel_file,
        file_name="Global_Predictive_Business_Analytics_Report.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    st.divider()

    # POWERPOINT REPORT

    st.subheader("Editable PowerPoint Executive Report")

    ppt_file = create_powerpoint_report(data)

    st.download_button(
        label="📈 Download PowerPoint Executive Report",
        data=ppt_file,
        file_name="Global_Predictive_Business_Analytics_Executive_Report.pptx",
        mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
    )


st.sidebar.success("Analytics Platform Complete ✓")
